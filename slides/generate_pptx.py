#!/usr/bin/env python3
"""Exporta la presentación Sempitecno a PowerPoint (.pptx)."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

import requests
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Colores corporativos
AZUL = RGBColor(0x0B, 0x25, 0x45)
VERDE = RGBColor(0x1A, 0xBC, 0x9C)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRIS = RGBColor(0x4B, 0x55, 0x63)
GRIS_CLARO = RGBColor(0xF3, 0xF4, 0xF6)
GRIS_LINEA = RGBColor(0xE5, 0xE7, 0xEB)
NARANJA_L2 = RGBColor(0xF5, 0x9E, 0x0B)

SCRIPT_DIR = Path(__file__).parent
LOGO_CACHE = SCRIPT_DIR / "assets" / "logos"
OUTPUT = SCRIPT_DIR / "Transformacion_Digital_Soporte_Inteligente_Sempitecno.pptx"

LOGO_URLS = {
    "zabbix": "https://www.vectorlogo.zone/logos/zabbix/zabbix-ar21.svg",
    "glpi": "https://raw.githubusercontent.com/glpi-project/glpi/main/public/pics/logos/logo-GLPI-250-black.png",
    "otobo": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/otrs.svg",
    "whatsapp": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/whatsapp.svg",
    "telegram": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/telegram.svg",
    "aws": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/amazonaws.svg",
    "azure": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/microsoftazure.svg",
    "gcp": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/googlecloud.svg",
    "openai": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/openai.svg",
    "issabel": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/asterisk.svg",
    "facebook": "https://cdn.jsdelivr.net/npm/simple-icons@11.14.0/icons/facebook.svg",
}


@dataclass
class SlideSpec:
    kind: str  # cover | section | content | cards | table | flow | contact
    title: str = ""
    subtitle: str = ""
    headline: str = ""
    bullets: list[str] = field(default_factory=list)
    cards: list[tuple[str, str]] = field(default_factory=list)
    logos: list[str] = field(default_factory=list)
    table_rows: list[tuple[str, str, str]] = field(default_factory=list)
    flow_steps: list[str] = field(default_factory=list)
    dark: bool = False
    badge: str = ""  # L1 | L2


class PptxBuilder:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self._logo_paths: dict[str, Path] = {}

    def _blank(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def _fill_bg(self, slide, color: RGBColor):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color

    def _watermark(self, slide, light: bool = False):
        box = slide.shapes.add_textbox(Inches(10.3), Inches(0.28), Inches(2.6), Inches(0.45))
        p = box.text_frame.paragraphs[0]
        r1 = p.add_run()
        r1.text = "Sempi"
        r1.font.bold = True
        r1.font.size = Pt(16)
        r1.font.color.rgb = BLANCO if light else AZUL
        r2 = p.add_run()
        r2.text = "tecno"
        r2.font.bold = True
        r2.font.size = Pt(16)
        r2.font.color.rgb = VERDE

    def _title_bar(self, slide, title: str, top=Inches(0.55), badge: str = ""):
        if badge:
            b = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), top, Inches(1.1), Inches(0.38)
            )
            b.fill.solid()
            b.fill.fore_color.rgb = VERDE if badge == "L1" else NARANJA_L2
            b.line.fill.background()
            bt = b.text_frame.paragraphs[0]
            bt.text = f"NIVEL {badge}"
            bt.font.size = Pt(11)
            bt.font.bold = True
            bt.font.color.rgb = BLANCO
            bt.alignment = PP_ALIGN.CENTER
            title_left = Inches(1.95)
        else:
            title_left = Inches(0.7)

        box = slide.shapes.add_textbox(title_left, top, Inches(11), Inches(0.75))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(30)
        p.font.color.rgb = BLANCO if False else AZUL
        slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
            Inches(0.7),
            top + Inches(0.72),
            Inches(2.4),
            Inches(0.055),
        ).fill.solid()
        slide.shapes[-1].fill.fore_color.rgb = VERDE
        slide.shapes[-1].line.fill.background()

    def _subtitle(self, slide, text: str, top=Inches(1.45)):
        if not text:
            return
        box = slide.shapes.add_textbox(Inches(0.7), top, Inches(11.5), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(15)
        p.font.color.rgb = GRIS

    def _bullets(self, slide, items: list[str], top=Inches(1.95)):
        box = slide.shapes.add_textbox(Inches(0.75), top, Inches(11.8), Inches(4.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"▸  {item}"
            p.font.size = Pt(17)
            p.font.color.rgb = GRIS
            p.space_after = Pt(8)

    def _ensure_logos(self):
        if self._logo_paths:
            return
        LOGO_CACHE.mkdir(parents=True, exist_ok=True)
        headers = {"User-Agent": "Mozilla/5.0"}
        for key, url in LOGO_URLS.items():
            path = LOGO_CACHE / f"{key}.png"
            if path.exists() and path.stat().st_size > 200:
                self._logo_paths[key] = path
                continue
            try:
                r = requests.get(url, headers=headers, timeout=20)
                r.raise_for_status()
                if b"<svg" in r.content[:500]:
                    import cairosvg
                    cairosvg.svg2png(bytestring=r.content, write_to=str(path), output_width=400)
                else:
                    img = Image.open(io.BytesIO(r.content)).convert("RGBA")
                    img.thumbnail((320, 140), Image.Resampling.LANCZOS)
                    img.save(path, "PNG")
                self._logo_paths[key] = path
            except Exception:
                pass

    def _add_logos(self, slide, keys: list[str], top=Inches(1.85)):
        self._ensure_logos()
        valid = [k for k in keys if k in self._logo_paths]
        if not valid:
            return
        n = len(valid)
        total_w = Inches(1.5) * n + Inches(0.3) * (n - 1)
        start = (self.prs.slide_width - total_w) / 2
        for i, key in enumerate(valid):
            left = start + i * Inches(1.8)
            slide.shapes.add_picture(str(self._logo_paths[key]), left, top, height=Inches(0.55))

    def _cards(self, slide, cards: list[tuple[str, str]], top=Inches(2.35)):
        n = len(cards)
        w = Inches(3.85)
        gap = Inches(0.35)
        total = n * w + (n - 1) * gap
        x0 = (self.prs.slide_width - total) / 2
        for i, (heading, body) in enumerate(cards):
            x = x0 + i * (w + gap)
            card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, top, w, Inches(3.6))
            card.fill.solid()
            card.fill.fore_color.rgb = GRIS_CLARO
            card.line.color.rgb = GRIS_LINEA
            bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, top, w, Inches(0.07))
            bar.fill.solid()
            bar.fill.fore_color.rgb = VERDE
            bar.line.fill.background()
            tb = slide.shapes.add_textbox(x + Inches(0.25), top + Inches(0.35), w - Inches(0.5), Inches(3.1))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = heading
            p.font.bold = True
            p.font.size = Pt(18)
            p.font.color.rgb = AZUL
            p2 = tf.add_paragraph()
            p2.text = body
            p2.font.size = Pt(13)
            p2.font.color.rgb = GRIS

    def _table(self, slide, rows: list[tuple[str, str, str]]):
        self._title_bar(slide, "L1 vs L2: Roles del Ingeniero IA")
        self._watermark(slide)
        cols, row_count = 3, len(rows) + 1
        left, top, width, height = Inches(0.65), Inches(1.75), Inches(12.0), Inches(4.8)
        table = slide.shapes.add_table(row_count, cols, left, top, width, height).table
        headers = ("Dimensión", "IA Nivel 1", "IA Nivel 2")
        hcolors = (AZUL, VERDE, NARANJA_L2)
        for c, (h, hc) in enumerate(zip(headers, hcolors)):
            cell = table.cell(0, c)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = hc
            for p in cell.text_frame.paragraphs:
                p.font.bold = True
                p.font.size = Pt(12)
                p.font.color.rgb = BLANCO
                p.alignment = PP_ALIGN.CENTER
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                cell = table.cell(r, c)
                cell.text = val
                if r % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = GRIS_CLARO
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(11)
                    p.font.color.rgb = AZUL if c == 0 else GRIS
                    p.font.bold = c == 0

    def _flow(self, slide, steps: list[str]):
        self._title_bar(slide, "Flujo: del usuario al especialista")
        self._subtitle(slide, "Trazabilidad completa en cada escalamiento.")
        self._watermark(slide)
        colors = [GRIS, VERDE, NARANJA_L2, AZUL]
        n = len(steps)
        box_w = Inches(2.7)
        gap = Inches(0.45)
        total = n * box_w + (n - 1) * gap
        x0 = (self.prs.slide_width - total) / 2
        y = Inches(2.35)
        for i, step in enumerate(steps):
            x = x0 + i * (box_w + gap)
            sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, box_w, Inches(1.05))
            sh.fill.solid()
            sh.fill.fore_color.rgb = colors[i % len(colors)]
            sh.line.fill.background()
            p = sh.text_frame.paragraphs[0]
            p.text = step
            p.alignment = PP_ALIGN.CENTER
            p.font.bold = True
            p.font.size = Pt(14)
            p.font.color.rgb = BLANCO
            if i < n - 1:
                arr = slide.shapes.add_shape(
                    MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, x + box_w + Inches(0.05), y + Inches(0.32), Inches(0.35), Inches(0.35)
                )
                arr.fill.solid()
                arr.fill.fore_color.rgb = GRIS
                arr.line.fill.background()
        self._bullets(
            slide,
            [
                "Registro en OTOBO/GLPI con historial y evidencia.",
                "Supervisores auditan decisiones del agente IA.",
                "Cada cierre alimenta la base de conocimiento.",
            ],
            top=Inches(3.85),
        )

    def add(self, spec: SlideSpec):
        slide = self._blank()

        if spec.kind == "cover":
            self._fill_bg(slide, AZUL)
            t = slide.shapes.add_textbox(Inches(0.85), Inches(2.2), Inches(11.5), Inches(2.2))
            p = t.text_frame.paragraphs[0]
            p.text = spec.title.replace("\\n", "\n")
            p.font.bold = True
            p.font.size = Pt(44)
            p.font.color.rgb = BLANCO
            s = slide.shapes.add_textbox(Inches(0.85), Inches(4.2), Inches(10), Inches(0.6))
            sp = s.text_frame.paragraphs[0]
            sp.text = spec.subtitle
            sp.font.size = Pt(22)
            sp.font.color.rgb = VERDE
            b = slide.shapes.add_textbox(Inches(0.85), Inches(5.5), Inches(4), Inches(0.5))
            bp = b.text_frame.paragraphs[0]
            r1 = bp.add_run()
            r1.text = "Sempi"
            r1.font.bold = True
            r1.font.size = Pt(26)
            r1.font.color.rgb = BLANCO
            r2 = bp.add_run()
            r2.text = "tecno"
            r2.font.bold = True
            r2.font.size = Pt(26)
            r2.font.color.rgb = VERDE
            return

        if spec.kind == "section":
            self._fill_bg(slide, AZUL)
            sec = slide.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(11), Inches(0.5))
            sec.text_frame.paragraphs[0].text = spec.subtitle
            sec.text_frame.paragraphs[0].font.size = Pt(18)
            sec.text_frame.paragraphs[0].font.color.rgb = VERDE
            tit = slide.shapes.add_textbox(Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.5))
            tp = tit.text_frame.paragraphs[0]
            tp.text = spec.title
            tp.font.bold = True
            tp.font.size = Pt(40)
            tp.font.color.rgb = BLANCO
            st = slide.shapes.add_textbox(Inches(0.9), Inches(4.5), Inches(11), Inches(0.6))
            st.text_frame.paragraphs[0].text = spec.headline
            st.text_frame.paragraphs[0].font.size = Pt(18)
            st.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
            return

        if spec.kind == "contact":
            self._fill_bg(slide, AZUL)
            tit = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(11), Inches(1.0))
            tit.text_frame.paragraphs[0].text = spec.title
            tit.text_frame.paragraphs[0].font.bold = True
            tit.text_frame.paragraphs[0].font.size = Pt(40)
            tit.text_frame.paragraphs[0].font.color.rgb = BLANCO
            sub = slide.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(10), Inches(0.5))
            sub.text_frame.paragraphs[0].text = spec.subtitle
            sub.text_frame.paragraphs[0].font.size = Pt(20)
            sub.text_frame.paragraphs[0].font.color.rgb = VERDE
            for i, line in enumerate(spec.bullets):
                box = slide.shapes.add_textbox(Inches(1.1), Inches(3.3) + i * Inches(0.55), Inches(6), Inches(0.5))
                p = box.text_frame.paragraphs[0]
                p.text = line
                p.font.size = Pt(20 if i == 0 else 17)
                p.font.bold = i == 0
                p.font.color.rgb = BLANCO if i == 0 else RGBColor(0xE2, 0xE8, 0xF0)
            return

        self._fill_bg(slide, BLANCO)

        if spec.kind == "table":
            self._table(slide, spec.table_rows)
            return
        if spec.kind == "flow":
            self._flow(slide, spec.flow_steps)
            return

        self._watermark(slide)
        self._title_bar(slide, spec.title, badge=spec.badge)
        self._subtitle(slide, spec.subtitle or spec.headline)

        if spec.logos:
            self._add_logos(slide, spec.logos, top=Inches(1.55) if spec.kind == "cards" else Inches(2.0))

        if spec.kind == "cards":
            self._cards(slide, spec.cards, top=Inches(2.55) if spec.logos else Inches(2.2))
        elif spec.kind == "pyramid":
            self._subtitle(slide, spec.headline, top=Inches(1.42))
            levels = [
                ("Nivel 3 — Especialistas humanos (cambios críticos)", AZUL, 0.5),
                ("Nivel 2 — Agente IA: diagnóstico, scripts, correlación", NARANJA_L2, 0.65),
                ("Nivel 1 — Agente IA: triaje y resolución rutinaria 24/7", VERDE, 0.8),
            ]
            y = Inches(2.15)
            for text, color, frac in levels:
                w = Inches(10) * frac
                left = (self.prs.slide_width - w) / 2
                bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, y, w, Inches(0.65))
                bar.fill.solid()
                bar.fill.fore_color.rgb = color
                bar.line.fill.background()
                p = bar.text_frame.paragraphs[0]
                p.text = text
                p.font.size = Pt(12)
                p.font.color.rgb = BLANCO
                p.alignment = PP_ALIGN.CENTER
                y += Inches(0.78)
            if spec.logos:
                self._add_logos(slide, spec.logos, top=Inches(5.0))
        elif spec.kind == "clients":
            clients = spec.bullets
            w, h = Inches(3.6), Inches(1.2)
            gap = Inches(0.3)
            x0 = Inches(0.75)
            y0 = Inches(2.3)
            for i, name in enumerate(clients):
                r, c = divmod(i, 3)
                x = x0 + c * (w + gap)
                y = y0 + r * (h + gap)
                box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
                box.fill.solid()
                box.fill.fore_color.rgb = GRIS_CLARO
                box.line.color.rgb = GRIS_LINEA
                box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = box.text_frame.paragraphs[0]
                p.text = name
                p.alignment = PP_ALIGN.CENTER
                p.font.bold = True
                p.font.size = Pt(20)
                p.font.color.rgb = AZUL
        elif spec.kind == "about":
            self._bullets(slide, spec.bullets, top=Inches(1.9))
            card = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(2.0), Inches(5.0), Inches(3.5)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = AZUL
            card.line.fill.background()
            tf = card.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.text = "20+"
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(64)
            p.font.bold = True
            p.font.color.rgb = VERDE
            p2 = tf.add_paragraph()
            p2.text = "años impulsando\ninfraestructura y soporte"
            p2.alignment = PP_ALIGN.CENTER
            p2.font.size = Pt(18)
            p2.font.color.rgb = BLANCO
        else:
            if spec.bullets:
                self._bullets(slide, spec.bullets, top=Inches(2.15) if spec.logos else Inches(1.95))

    def save(self, path: Path):
        self.prs.save(path)


def build_specs() -> list[SlideSpec]:
    return [
        SlideSpec("cover", title="Transformación Digital\ny Soporte Inteligente", subtitle="Soluciones de IT que impulsan su negocio"),
        SlideSpec("about", title="Sobre Nosotros", bullets=[
            "Más de 20 años de experiencia en el mercado.",
            "Entregamos el código fuente de soluciones tipo open-source.",
            "Capacitación exhaustiva a su equipo.",
            "Asociación transparente en su crecimiento tecnológico.",
        ]),
        SlideSpec("clients", title="Clientes que confían en nosotros", subtitle="Empresas líderes que respaldan la calidad de nuestras soluciones.", bullets=[
            "Telefónica", "Fybeca", "Veris", "Pycca", "Médikal", "Bankers Club",
        ]),
        SlideSpec("section", subtitle="Sección 1", title="Plataformas Robustas de Infraestructura", headline="Control, Privacidad y Estabilidad Operativa"),
        SlideSpec("content", title="Zabbix: Monitoreo Total", headline="El guardián de su red", logos=["zabbix"], bullets=[
            "Monitorea servidores, switches, impresoras y UPS.",
            "Detecta anomalías antes de que afecten a los usuarios.",
            "Notifica incidentes en tiempo real.",
        ]),
        SlideSpec("cards", title="Zabbix: Casos de Uso", logos=["zabbix"], cards=[
            ("Autorrecuperación", "Reinicio automático de servicios críticos sin intervención humana."),
            ("Bases de Datos", "Monitoreo profundo del rendimiento para respuestas rápidas."),
            ("Dispositivos Físicos", "IoT, antenas y cámaras bajo control de disponibilidad."),
        ]),
        SlideSpec("content", title="GLPI: Gestión de Activos", headline="Control total de lo que tiene", logos=["glpi"], bullets=[
            "Inventario automático de hardware y software.",
            "Licencias, contratos y garantías.",
            "Proveedores y ciclo de vida del equipo.",
        ]),
        SlideSpec("cards", title="GLPI: Casos de Uso", logos=["glpi"], cards=[
            ("Proyectos IT", "Organización y seguimiento de inicio a cierre."),
            ("Reportes Financieros", "Auditorías, depreciación y presupuestos."),
            ("Mesa de Ayuda", "Incidencias vinculadas a activos afectados."),
        ]),
        SlideSpec("content", title="OTOBO: Mesa de Servicio", headline="Atención estructurada y eficiente", logos=["otobo"], bullets=[
            "Workflows totalmente personalizados.",
            "Certificado en mejores prácticas ITIL.",
            "Correos transformados automáticamente en tickets.",
        ]),
        SlideSpec("cards", title="OTOBO: Casos de Uso", logos=["otobo"], cards=[
            ("Seguimiento Veloz", "SLA visibles en tiempo real."),
            ("Búsquedas Avanzadas", "Historial de soluciones para evitar retrabajo."),
            ("Base de Conocimientos", "Portal FAQ de autoservicio."),
        ]),
        SlideSpec("content", title="Issabel: Telefonía IP", headline="Comunicaciones sin límites", logos=["issabel"], bullets=[
            "Sin costo de licencias por usuario.",
            "Extensiones telefónicas ilimitadas.",
            "IVR y operadora automática incluidos.",
        ]),
        SlideSpec("cards", title="Issabel: Casos de Uso", logos=["issabel"], cards=[
            ("Grabación", "Auditoría y control de calidad de llamadas."),
            ("Grupos de Marcado", "Distribución por departamento."),
            ("Voz a Medida", "Menús, buzones a correo y conferencias."),
        ]),
        SlideSpec("section", subtitle="Sección 2", title="Automatización con IA y Soporte Digital", headline="Ingenieros de soporte L1 y L2 disponibles 24/7"),
        SlideSpec("pyramid", title="Agentes de IA como Ingenieros de Soporte", headline="Ingenieros digitales L1 y L2 integrados a su stack — no bots de respuestas fijas.", logos=["openai"]),
        SlideSpec("content", title="Ingeniero de Soporte IA", subtitle="Primera línea digital: disponible 24/7 sin colas de espera.", badge="L1", bullets=[
            "Atiende como analista de mesa: saluda, clasifica y prioriza.",
            "Resuelve rutina: contraseñas, VPN, impresoras, accesos.",
            "Consulta base de conocimiento (OTOBO/GLPI) paso a paso.",
            "Crea tickets con categoría, activo y pasos ya ejecutados.",
            "Deriva a L2 solo si detecta complejidad o falta de permisos.",
        ], logos=["whatsapp", "telegram", "openai"]),
        SlideSpec("content", title="Ingeniero de Soporte IA", subtitle="Segunda línea digital: diagnóstico técnico y remediación controlada.", badge="L2", bullets=[
            "Se conecta a Zabbix: correlaciona alertas y tendencias.",
            "Ejecuta comandos autorizados (SSH/PowerShell/API).",
            "Aplica runbooks: reinicios, colas, verificación de backups.",
            "Enriquece tickets en GLPI/OTOBO con evidencia y causa raíz sugerida.",
            "Escala a L3 con timeline, comandos e impacto documentado.",
        ], logos=["zabbix", "glpi", "otobo"]),
        SlideSpec("table", table_rows=[
            ("Objetivo", "Desvío de volumen y autoservicio", "Diagnóstico y remediación"),
            ("Usuario", "Colaborador / cliente final", "TI interna / operaciones"),
            ("Acciones", "Guías, FAQs, tickets simples", "Scripts, APIs, runbooks"),
            ("Integraciones", "OTOBO, chat, correo", "Zabbix, GLPI, servidores"),
            ("Escalamiento", "→ L2 IA o humano", "→ L3 especialista"),
        ]),
        SlideSpec("flow", flow_steps=["Usuario", "IA L1", "IA L2", "Humano L3"]),
        SlideSpec("cards", title="Agente Técnico: Integración Avanzada", logos=["zabbix", "openai", "glpi"], cards=[
            ("Diagnóstico preventivo", "Consulta Zabbix antes del reporte humano."),
            ("Ejecución controlada", "SSH/PowerShell con políticas y auditoría."),
            ("Escalamiento inteligente", "Tickets enriquecidos en GLPI/OTOBO."),
        ]),
        SlideSpec("content", title="Chatbots para WhatsApp y Web", headline="Atención inmediata en canales favoritos", logos=["whatsapp", "telegram"], bullets=[
            "Integración nativa con WhatsApp, Web y redes sociales.",
            "Conexión directa a CRM o ERP corporativo.",
            "Botones interactivos y listas rápidas.",
        ]),
        SlideSpec("cards", title="Chatbots: Casos de Uso", cards=[
            ("Venta automatizada", "Catálogos y pedidos desde el chat."),
            ("Encuestas NPS", "Satisfacción post-atención automática."),
            ("Servicios internos", "Vacaciones, roles de pago, reseteo de claves."),
        ]),
        SlideSpec("content", title="Agente Virtual Cognitivo de Voz", logos=["openai", "issabel"], bullets=[
            "Voz humana realista, sin tonos robóticos.",
            "Entiende contexto conversacional fluido.",
            "Analiza y responde en tiempo real.",
        ]),
        SlideSpec("cards", title="Voice Bots: Casos de Uso", cards=[
            ("Agendamiento", "Citas médicas, mecánicas o de asesoría."),
            ("Llamadas salientes", "Recordatorios y confirmaciones."),
            ("Encuestas telefónicas", "Sondeos conversacionales de calidad."),
        ]),
        SlideSpec("content", title="Chat Multiagente y Multicanal", logos=["whatsapp", "telegram", "facebook"], bullets=[
            "Un buzón para WhatsApp, Telegram, Messenger y correo.",
            "Múltiples agentes en el mismo número.",
            "Supervisión, auditoría y asignación de chats.",
        ]),
        SlideSpec("cards", title="Multicanalidad: Casos de Uso", cards=[
            ("Atención masiva", "Cientos de clientes sin mezclar historiales."),
            ("Soporte organizado", "Transferencia L1→L2 con contexto intacto."),
            ("Métricas", "Tiempos de respuesta por canal y agente."),
        ]),
        SlideSpec("cards", title="Beneficios para su Empresa", cards=[
            ("Downtime reducido", "Monitoreo proactivo y alertas tempranas."),
            ("Resolución instantánea", "IA L1/L2 atiende rutina y acorta MTTR."),
            ("Talento liberado", "Humanos en innovación, no en reseteos."),
        ]),
        SlideSpec("content", title="Nuestra Propuesta de Valor", logos=["aws", "azure", "gcp"], bullets=[
            "Soluciones llave en mano: implementamos, capacitamos y entregamos listo.",
            "Sin vendor lock-in: modelos abiertos y código fuente.",
            "Flexibilidad: On-Premise o nube (AWS, Azure, Google Cloud).",
        ]),
        SlideSpec("contact", title="¿Listo para dar el siguiente paso?", subtitle="Hablemos sobre cómo optimizar su negocio.", bullets=[
            "Iván García", "+593 999989522", "ivang@sempitecno.com", "www.sempitecno.com",
        ]),
    ]


def main():
    builder = PptxBuilder()
    for spec in build_specs():
        builder.add(spec)
    builder.save(OUTPUT)
    print(f"PPTX generado: {OUTPUT}")
    print(f"Tamaño: {OUTPUT.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
